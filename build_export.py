"""Portable XLSX export. Only the caller's selected local media are embedded."""
import json
import sys
from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

def literal(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value

def build(data, output):
    wb = Workbook()
    wb.remove(wb.active)
    keep_images = []
    def sheet(name, title, headers, widths):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws["A2"] = title
        ws["A2"].font = Font(name="Arial", size=15, bold=True)
        ws["A3"] = "来源、缺项、款龄依据与审核结论分开记录；空白不代表0。"
        ws.freeze_panes = "A6"
        for col, (header, width) in enumerate(zip(headers, widths), 1):
            c = ws.cell(5, col, header)
            c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="504466")
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[5].height = 30
        return ws
    def row(ws, idx, values):
        for col, value in enumerate(values, 1):
            c = ws.cell(idx, col, literal(value))
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[idx].height = 150
    def image(ws, source, anchor, width=120, height=175):
        if not source: return
        try:
            im = Image(str(source))
            ratio = min(width / im.width, height / im.height)
            im.width *= ratio
            im.height *= ratio
            ws.add_image(im, anchor)
            keep_images.append(im)
        except Exception:
            if data.get("strict_images"): raise
            ws[anchor] = "图片文件待补"
    headers = ["店铺名称","三级类目","商品ID / SKU","正面图","背面图","原销量标签","款龄结论","首次上架日","颜色选项数","具体颜色","审核结论","审核人","缺项 / 理由","商品来源","日期证据","主风格","元素提取","销量周期","SKU月份线索","来源任务","渠道与筛选规则","审核版本"]
    main = sheet("选款图文","选款审核与图文交付",headers,[16,20,26,20,20,16,24,16,12,26,20,18,40,44,36,20,40,24,20,30,46,24])
    colors = sheet("复色图片","各颜色选项原图",["商品ID","三级类目","选项序号","平台颜色名","颜色图片","原图来源","来源任务"],[24,24,14,24,22,60,30])
    color_row = 6
    for idx, item in enumerate(data.get("rows", []), 6):
        rule = item.get("source_rules") or data.get("rules", {})
        review = item.get("review") or {}
        evaluation = item.get("evaluation") or {}
        first_date = review.get("first_listed_at") or item.get("first_listed_at")
        try: first_date = date.fromisoformat(first_date) if first_date else None
        except ValueError: pass
        leaf = item.get("category_leaf") or (item.get("breadcrumb") or ["待核"])[-1]
        status = "审核已确认" if evaluation.get("approved") else {"eligible":"符合规则","pending":"待核","excluded":"规则排除","rejected":"人工排除"}.get(evaluation.get("status"),"待核")
        vals = [rule.get("store"),leaf,str(item.get("id",""))+" / "+str(item.get("sku") or "待核"),None if item.get("front") else "待补",None if item.get("back") else "待补",item.get("sales"),evaluation.get("age_status"),first_date,item.get("color_count"),
            "\n".join(f"{i+1}. {c.get('name','')}" for i,c in enumerate(item.get("colors",[]))),status,review.get("reviewer"),
            "；".join(evaluation.get("reasons",[]))+("；"+review["reason"] if review.get("reason") else ""),item.get("url"),review.get("date_source") or item.get("date_source") or "待核",
            review.get("primary_style") or "待图审",review.get("elements") or "待图审",evaluation.get("sales_period") or "平台未公开周期",evaluation.get("sku_month_clue") or "无",item.get("job") or data.get("job",""),
            f"{rule.get('platform','')} / {rule.get('country','')}；销量≥{rule.get('min_sold','')}；{rule.get('max_age_months','')}个月；截至{rule.get('as_of','')}；{rule.get('age_basis','')}",item.get("revision")]
        row(main,idx,vals)
        main.cell(idx,8).number_format = "yyyy-mm-dd"
        image(main,item.get("front"),f"D{idx}")
        image(main,item.get("back"),f"E{idx}")
        for i,c in enumerate(item.get("colors",[]),1):
            row(colors,color_row,[item.get("id"),leaf,i,c.get("name"),None if c.get("image") else "待补",c.get("source_url") or c.get("src") or item.get("url"),item.get("job") or data.get("job")])
            image(colors,c.get("image"),f"E{color_row}")
            color_row += 1
    if not data.get("rows"):main["A6"]="当前范围无记录；请补齐证据和人工审核。"
    if color_row == 6:colors["A6"]="当前范围暂无颜色记录"
    if data.get("designs"):
        ws = sheet("改款参考","改款方案与AI概念参考图",["源商品ID","方案名称","保留内容","改动内容","概念参考图","生成 / 审核状态","提示词"],[20,28,38,42,30,30,60])
        for idx,design in enumerate(data["designs"],6):
            pics=design.get("images") or []
            status="设计师已确认" if (design.get("review") or {}).get("decision")=="approved" else "已生成，待设计师确认" if pics else "方案草案"
            row(ws,idx,[design.get("product_id"),design.get("name"),design.get("keep"),design.get("changes"),None if pics else "尚未生成",status,design.get("prompt")])
            if pics:image(ws,pics[0].get("path"),f"E{idx}")
    for ws in wb:
        ws.auto_filter.ref=f"A5:{get_column_letter(ws.max_column)}{max(5,ws.max_row)}"
    out=Path(output);out.parent.mkdir(parents=True,exist_ok=True);wb.save(out)
    return {"saved":str(out),"rows":len(data.get("rows",[])),"embedded_images":len(keep_images)}
if __name__=="__main__":
    print(json.dumps(build(json.loads(Path(sys.argv[1]).read_text()),sys.argv[2]),ensure_ascii=False))
